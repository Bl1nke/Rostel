"""Экспорт расчётов калькулятора в DOCX и CSV.

Модуль не выполняет вычислений: получает готовый результат ``calculate()``
и формирует документ или таблицу. Не пересчитывает стоимость и не расширяет
контракт ``calculate()`` — только читает поля из результата.
"""
from __future__ import annotations

import csv
from datetime import date, datetime
from io import BytesIO, StringIO
from pathlib import Path
from typing import Any, Iterable

try:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    _DOCX_AVAILABLE = True
except ImportError:  # pragma: no cover - окружение без python-docx
    _DOCX_AVAILABLE = False
    Document = None  # type: ignore[assignment]
    WD_TABLE_ALIGNMENT = None  # type: ignore[assignment]
    WD_ALIGN_PARAGRAPH = None  # type: ignore[assignment]
    Pt = None  # type: ignore[assignment]
    RGBColor = None  # type: ignore[assignment]


# ============================================================================
# Константы
# ============================================================================

PROPOSAL_TITLE = "Коммерческое предложение"

PROPOSAL_DISCLAIMER = (
    "Документ сформирован автоматически и содержит предварительную оценку "
    "стоимости. Он не является публичной офертой и не влечёт финансовых или "
    "юридических обязательств. Итоговая стоимость уточняется после "
    "согласования технического задания."
)

CALCULATION_DISCLAIMER = (
    "Стоимость рассчитана как произведение базовой ставки за м² (зависит от "
    "типа объекта) на площадь и поправочные коэффициенты: региональный "
    "(инвестиционный климат), срочность (по дате завершения), схема "
    "финансирования, вид работ, тип оборудования и наличие уникальных "
    "технических решений."
)

MARGIN_DISCLAIMER = (
    "Маржинальность формируется как базовая по виду работ (20–25 %) плюс "
    "поправка на конкуренцию (число открытых тендеров в регионе и "
    "направлении). Прибыль = цена − себестоимость; "
    "маржинальность = прибыль / цена × 100 %."
)

ANALOGS_DISCLAIMER = (
    "Сопоставимые проекты отбираются из исторических данных по совпадению "
    "региона и близости суммы к рассчитанному диапазону цен. Если в целевом "
    "регионе совпадений нет, показываются проекты схожего масштаба из других "
    "регионов — это явно отмечено в таблице."
)

CSV_DELIMITER = ";"
CSV_ENCODING = "utf-8-sig"
CSV_LINE_TERMINATOR = "\r\n"

CSV_HEADER: tuple[str, ...] = (
    "Дата расчёта",
    "Тип объекта",
    "Площадь, м²",
    "Регион ID",
    "Регион",
    "Дедлайн",
    "Срочность",
    "Коэффициент срочности",
    "Финансирование",
    "Вид работ",
    "Оборудование",
    "Уникальные решения",
    "Ставка за м², ₽",
    "Мин. цена, ₽",
    "Оптимальная цена, ₽",
    "Макс. цена, ₽",
    "Себестоимость, ₽",
    "Прибыль, ₽",
    "Маржинальность, %",
    "Кол-во аналогов",
)

_EMPTY = "—"
_ACCENT_RGB: tuple[int, int, int] = (0xFF, 0x4F, 0x12)  # фирменный оранжевый
_MUTED_RGB: tuple[int, int, int] = (0x6B, 0x6B, 0x80)


# ============================================================================
# Форматирование значений
# ============================================================================

def _format_rub(value: Any) -> str:
    """Форматирует сумму в рублях с пробелами-разделителями тысяч."""
    if value is None:
        return _EMPTY
    try:
        number = float(value)
    except (TypeError, ValueError):
        return _EMPTY
    return f"{int(round(number)):,}".replace(",", " ") + " ₽"


def _format_percent(value: Any) -> str:
    """Форматирует процент с одним знаком после запятой."""
    if value is None:
        return _EMPTY
    try:
        number = float(value)
    except (TypeError, ValueError):
        return _EMPTY
    return f"{number:.1f}%".replace(".", ",")


def _format_int(value: Any) -> str:
    """Форматирует целое число с пробелами как разделителями тысяч."""
    if value is None:
        return _EMPTY
    try:
        number = int(round(float(value)))
    except (TypeError, ValueError):
        return _EMPTY
    return f"{number:,}".replace(",", " ")


def _normalize_date(value: date | datetime | str | None) -> str:
    """Возвращает дату в ISO-формате (YYYY-MM-DD)."""
    if value is None:
        return date.today().isoformat()
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def _format_date_ru(value: date | datetime | str | None) -> str:
    """Возвращает дату в формате DD.MM.YYYY (русский стандарт)."""
    iso = _normalize_date(value)
    try:
        return date.fromisoformat(iso).strftime("%d.%m.%Y")
    except (TypeError, ValueError):
        return iso


def _csv_number(value: Any, decimals: int = 2) -> str:
    """Число для CSV: запятая как десятичный разделитель, пусто при None."""
    if value is None or value == "":
        return ""
    try:
        number = float(value)
    except (TypeError, ValueError):
        return ""
    if decimals == 0:
        return str(int(round(number)))
    return f"{number:.{decimals}f}".replace(".", ",")


def _csv_number_auto(value: Any) -> str:
    """Целое — без дробной части, дробное — с 2 знаками, запятая."""
    if value is None or value == "":
        return ""
    try:
        number = float(value)
    except (TypeError, ValueError):
        return ""
    if number == int(number):
        return str(int(number))
    return f"{number:.2f}".replace(".", ",")


def _csv_text(value: Any) -> str:
    """Текст для CSV без ``None`` и лишних пробелов."""
    if value is None:
        return ""
    return str(value).strip()


def _join_region(region_name: Any, region_id: Any) -> str:
    """Собирает читаемое название региона из имени и/или ID."""
    name = _csv_text(region_name)
    rid = _csv_text(region_id)
    if name and rid and name != rid:
        return f"{name} ({rid})"
    return name or rid or _EMPTY


def _profit_of(result: dict[str, Any]) -> float | None:
    """Возвращает прибыль: из результата, иначе вычисляет optimal − cost."""
    profit = result.get("profit")
    if profit is not None:
        try:
            return float(profit)
        except (TypeError, ValueError):
            return None
    optimal = result.get("optimal_price")
    cost = result.get("cost_of_goods_sold")
    if optimal is None or cost is None:
        return None
    try:
        return float(optimal) - float(cost)
    except (TypeError, ValueError):
        return None


def _urgency_text(params: dict[str, Any]) -> str:
    """Собирает строку вида «Повышенная — 60 дн. до завершения (×1.10)»."""
    label = _csv_text(params.get("urgency"))
    coefficient = params.get("urgency_coefficient")
    if not label and coefficient is None:
        return _EMPTY
    suffix = ""
    if coefficient is not None:
        try:
            suffix = f" (×{float(coefficient):.2f})"
        except (TypeError, ValueError):
            suffix = ""
    return f"{label or _EMPTY}{suffix}"


# ============================================================================
# DOCX — коммерческое предложение
# ============================================================================

def build_commercial_proposal_docx(
    result: dict[str, Any],
    *,
    proposal_date: date | datetime | str | None = None,
    customer_name: str | None = None,
) -> bytes:
    """Формирует коммерческое предложение в формате DOCX.

    Args:
        result: результат ``calculator.calculate()`` — не пересчитывается.
        proposal_date: дата документа; по умолчанию — сегодня.
        customer_name: необязательное имя заказчика. Если ``None``, строка
            «Заказчик» в документ не добавляется. Никакие другие реквизиты
            не подставляются.

    Raises:
        RuntimeError: библиотека ``python-docx`` не установлена.
        TypeError: ``result`` — не словарь.
        ValueError: в ``result`` нет ключа ``optimal_price``.
    """
    if not _DOCX_AVAILABLE:
        raise RuntimeError(
            "Библиотека python-docx недоступна. Установите её: "
            "pip install python-docx"
        )
    if not isinstance(result, dict):
        raise TypeError(
            "result должен быть результатом calculator.calculate() — словарём."
        )
    if "optimal_price" not in result:
        raise ValueError(
            "result не содержит ключа 'optimal_price'. Передайте результат "
            "calculator.calculate()."
        )

    document = Document()
    _setup_document(document)
    _add_title(document, proposal_date, customer_name)
    _add_object_section(document, result)
    _add_price_section(document, result)
    _add_cost_margin_section(document, result)
    _add_calculation_explanation(document, result)
    _add_analogs_section(document, result)
    _add_insight_section(document, result)
    _add_disclaimer(document)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def save_commercial_proposal_docx(
    result: dict[str, Any],
    path: str | Path,
    **kwargs: Any,
) -> Path:
    """Сохраняет DOCX по указанному пути и возвращает ``Path``."""
    data = build_commercial_proposal_docx(result, **kwargs)
    target = Path(path)
    target.write_bytes(data)
    return target


# ---------- Внутренние помощники DOCX ----------

def _setup_document(document: Any) -> None:
    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(4)


def _add_title(
    document: Any,
    proposal_date: date | datetime | str | None,
    customer_name: str | None,
) -> None:
    document.add_heading(PROPOSAL_TITLE, level=0)

    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = meta.add_run(f"Дата формирования: {_format_date_ru(proposal_date)}")
    run.italic = True
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(*_MUTED_RGB)

    if customer_name:
        line = document.add_paragraph()
        label = line.add_run("Заказчик: ")
        label.bold = True
        line.add_run(str(customer_name).strip())


def _add_object_section(document: Any, result: dict[str, Any]) -> None:
    params = result.get("project_params") or {}
    region_text = _join_region(result.get("region_name"), params.get("region_id"))

    rows: list[tuple[str, str]] = [
        ("Тип объекта", _csv_text(params.get("object_type")) or _EMPTY),
        ("Площадь, м²", _format_int(params.get("area_sqm"))),
        ("Регион", region_text),
        (
            "Дата завершения",
            _format_date_ru(params.get("deadline_date"))
            if params.get("deadline_date")
            else _EMPTY,
        ),
        ("Срочность", _urgency_text(params)),
        ("Финансирование", _csv_text(params.get("financing")) or _EMPTY),
        ("Вид работ", _csv_text(params.get("work_type")) or "не указан"),
        ("Оборудование", _csv_text(params.get("equipment")) or _EMPTY),
        ("Уникальные решения", "да" if params.get("unique_solution") else "нет"),
    ]

    base_rate = result.get("base_rate_per_sqm")
    if base_rate is not None:
        rows.insert(1, ("Базовая ставка за м²", _format_rub(base_rate)))

    document.add_heading("Параметры объекта", level=1)
    _add_kv_table(document, rows)


def _add_price_section(document: Any, result: dict[str, Any]) -> None:
    rows = [
        ("Минимальная стоимость", _format_rub(result.get("min_price"))),
        ("Оптимальная стоимость", _format_rub(result.get("optimal_price"))),
        ("Максимальная стоимость", _format_rub(result.get("max_price"))),
    ]
    document.add_heading("Стоимость проекта", level=1)
    _add_kv_table(document, rows, highlight_labels=("Оптимальная стоимость",))


def _add_cost_margin_section(document: Any, result: dict[str, Any]) -> None:
    cost = result.get("cost_of_goods_sold")
    margin = result.get("margin_percent")
    optimal = result.get("optimal_price")
    profit = _profit_of(result)

    document.add_heading("Себестоимость и маржинальность", level=1)

    if cost is None and margin is None:
        document.add_paragraph(
            "Себестоимость и маржинальность не рассчитаны: недостаточно "
            "достоверных данных в источнике."
        )
        return

    rows: list[tuple[str, str]] = []
    if cost is not None:
        rows.append(("Себестоимость", _format_rub(cost)))
    if margin is not None and optimal is not None and cost is not None:
        profit_value = profit if profit is not None else float(optimal) - float(cost)
        rows.append(
            (
                "Маржинальность",
                f"{_format_percent(margin)} — это {_format_rub(profit_value)} "
                f"при цене {_format_rub(optimal)}",
            )
        )
    elif margin is not None:
        rows.append(("Маржинальность", _format_percent(margin)))
    _add_kv_table(document, rows)
    document.add_paragraph(MARGIN_DISCLAIMER)


def _add_calculation_explanation(document: Any, result: dict[str, Any]) -> None:
    document.add_heading("Принцип расчёта", level=1)
    document.add_paragraph(CALCULATION_DISCLAIMER)

    coefficient = result.get("region_coefficient")
    if coefficient is None:
        return

    params = result.get("project_params") or {}
    region_text = _join_region(result.get("region_name"), params.get("region_id"))

    line = document.add_paragraph()
    label = line.add_run("Региональный коэффициент: ")
    label.bold = True
    line.add_run(f"{float(coefficient):.2f}".replace(".", ","))
    line.add_run(f" — инвестиционный климат региона «{region_text}».")


def _add_analogs_section(document: Any, result: dict[str, Any]) -> None:
    document.add_heading("Сопоставимые проекты", level=1)

    analog_count = result.get("analog_count")
    analog_data = result.get("analog_data") or []

    if analog_count is None:
        document.add_paragraph(
            "Исторические данные для сравнения недоступны: файл "
            "data/investmensts.csv не найден или не содержит обязательных "
            "колонок."
        )
        return
    if analog_count == 0 or not analog_data:
        document.add_paragraph(
            "Сопоставимых проектов в исторических данных не найдено."
        )
        return

    document.add_paragraph(
        f"Использовано проектов: {analog_count}. "
        "Ниже перечислены ближайшие по сумме к расчётному диапазону."
    )

    table = document.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for index, header in enumerate(("ID", "Проект / компания", "Регион", "Сумма", "ROI")):
        _set_cell_text(table.rows[0].cells[index], header, bold=True)

    for item in analog_data:
        cells = table.add_row().cells
        project = _csv_text(item.get("project_name")) or _EMPTY
        company = _csv_text(item.get("company_name"))
        combined = f"{project} / {company}" if company else project
        _set_cell_text(cells[0], _csv_text(item.get("investment_id")) or _EMPTY)
        _set_cell_text(cells[1], combined)
        region_text = _csv_text(item.get("region")) or _EMPTY
        if item.get("same_region") is False:
            region_text = f"{region_text} (другой регион)"
        _set_cell_text(cells[2], region_text)
        _set_cell_text(cells[3], _format_rub(item.get("amount_rub")))
        _set_cell_text(cells[4], _format_percent(item.get("expected_roi_percent")))

    document.add_paragraph(ANALOGS_DISCLAIMER)


def _add_insight_section(document: Any, result: dict[str, Any]) -> None:
    insight = result.get("business_insight")
    if not insight:
        return
    document.add_heading("Бизнес-инсайт", level=1)
    document.add_paragraph(str(insight))


def _add_disclaimer(document: Any) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(PROPOSAL_DISCLAIMER)
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(*_MUTED_RGB)


def _add_kv_table(
    document: Any,
    rows: list[tuple[str, str]],
    *,
    highlight_labels: Iterable[str] = (),
) -> None:
    """Двухколоночная таблица «label — value» с выделением избранных строк."""
    highlight = set(highlight_labels)
    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT

    for label, value in rows:
        cells = table.add_row().cells
        _set_cell_text(cells[0], label, bold=True)
        if label in highlight:
            _set_cell_text(cells[1], value, bold=True, color_rgb=_ACCENT_RGB)
        else:
            _set_cell_text(cells[1], value)


def _set_cell_text(
    cell: Any,
    text: Any,
    *,
    bold: bool = False,
    italic: bool = False,
    color_rgb: tuple[int, int, int] | None = None,
) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(str(text))
    if bold:
        run.bold = True
    if italic:
        run.italic = True
    if color_rgb is not None:
        run.font.color.rgb = RGBColor(*color_rgb)


# ============================================================================
# CSV — история расчётов
# ============================================================================

def build_calculation_history_csv(
    history: Iterable[dict[str, Any]] | dict[str, Any] | None,
    *,
    calc_date: date | datetime | str | None = None,
    delimiter: str = CSV_DELIMITER,
    encoding: str = CSV_ENCODING,
) -> bytes:
    """Формирует CSV с историей расчётов и возвращает байты.

    Каждый элемент ``history`` — результат ``calculator.calculate()``,
    опционально с ключом ``"calculated_at"`` (ISO-строка / date / datetime).
    Если ключа нет — используется ``calc_date`` или текущая дата.

    Формат вывода ориентирован на русскую локаль Excel:
        * разделитель ``;``;
        * десятичный разделитель — запятая;
        * кодировка ``utf-8-sig`` (BOM для корректного открытия в Excel);
        * перенос строк ``\\r\\n``;
        * отсутствующие значения — пустая строка.

    Одиночный словарь принимается за список из одного расчёта.
    """
    if history is None:
        items: list[dict[str, Any]] = []
    elif isinstance(history, dict):
        items = [history]
    else:
        items = [item for item in history if isinstance(item, dict)]

    default_date = _normalize_date(calc_date)

    buffer = StringIO()
    writer = csv.writer(
        buffer,
        delimiter=delimiter,
        quoting=csv.QUOTE_MINIMAL,
        lineterminator=CSV_LINE_TERMINATOR,
    )
    writer.writerow(CSV_HEADER)
    for item in items:
        writer.writerow(_history_row(item, default_date))

    return buffer.getvalue().encode(encoding)


def save_calculation_history_csv(
    history: Iterable[dict[str, Any]] | dict[str, Any] | None,
    path: str | Path,
    **kwargs: Any,
) -> Path:
    """Сохраняет CSV по указанному пути и возвращает ``Path``."""
    data = build_calculation_history_csv(history, **kwargs)
    target = Path(path)
    target.write_bytes(data)
    return target


def _history_row(result: dict[str, Any], default_date: str) -> tuple[str, ...]:
    params = result.get("project_params") or {}
    profit = _profit_of(result)
    return (
        _normalize_date(result.get("calculated_at") or default_date),
        _csv_text(params.get("object_type")),
        _csv_number_auto(params.get("area_sqm")),
        _csv_text(params.get("region_id")),
        _csv_text(result.get("region_name") or params.get("region_id")),
        _csv_text(params.get("deadline_date")),
        _csv_text(params.get("urgency")),
        _csv_number(params.get("urgency_coefficient"), 2),
        _csv_text(params.get("financing")),
        _csv_text(params.get("work_type")),
        _csv_text(params.get("equipment")),
        "да" if params.get("unique_solution") else "нет",
        _csv_number_auto(result.get("base_rate_per_sqm")),
        _csv_number(result.get("min_price"), 2),
        _csv_number(result.get("optimal_price"), 2),
        _csv_number(result.get("max_price"), 2),
        _csv_number(result.get("cost_of_goods_sold"), 2),
        _csv_number(profit, 2),
        _csv_number(result.get("margin_percent"), 1),
        _csv_number(result.get("analog_count"), 0),
    )


__all__ = [
    "PROPOSAL_TITLE",
    "PROPOSAL_DISCLAIMER",
    "CALCULATION_DISCLAIMER",
    "MARGIN_DISCLAIMER",
    "ANALOGS_DISCLAIMER",
    "CSV_HEADER",
    "build_commercial_proposal_docx",
    "build_calculation_history_csv",
    "save_commercial_proposal_docx",
    "save_calculation_history_csv",
]